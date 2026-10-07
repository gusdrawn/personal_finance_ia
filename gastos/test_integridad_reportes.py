from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from banca.models import CategoriaTematica, ImportacionEstadoCuenta, MovimientoBancario, SaldoSnapshot
from banca.services.presupuesto import promedio_ingresos, proyeccion_no_facturados, resumen_presupuesto, detalle_promedios_categoria
from configuracion.models import Banco, Producto, TipoCambio
from gastos.models import ConciliacionMovimientoManual, MovimientoManual
from gastos.services import monto_facturado_clp, resumen_gastos


class IntegridadReportesTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user('integridad_reportes')
        self.other = User.objects.create_user('integridad_otra')
        self.bank = Banco.objects.create(nombre='Banco reportes')
        self.account = Producto.objects.create(user=self.user, banco=self.bank, tipo='CUENTA', nombre='Cuenta')
        self.card = Producto.objects.create(user=self.user, banco=self.bank, tipo='TDC', nombre='Tarjeta')
        self.cat = CategoriaTematica.objects.create(user=self.user, nombre='Compras', aplica_a='AMBOS', esencialidad_defecto='OCIO')
        self.period = date(2026, 8, 1)

    def movimiento(self, amount='100', **kwargs):
        data = dict(user=self.user, producto=self.account, fecha=date(2026, 8, 10), monto=Decimal(amount), monto_clp=Decimal(amount), estado='FACTURADO', descripcion='Compra', categoria=self.cat, hash_externo=f'm-{MovimientoBancario.objects.count()}')
        data.update(kwargs)
        return MovimientoBancario.objects.create(**data)

    def estado(self, **kwargs):
        data = dict(user=self.user, producto=self.card, nombre_archivo='estado.pdf', hash_archivo=f'e-{ImportacionEstadoCuenta.objects.count()}', estado='PROCESADO', fecha_estado=date(2026, 7, 31), fecha_vencimiento=date(2026, 8, 10), monto_facturado=Decimal('100'), moneda='CLP')
        data.update(kwargs)
        return ImportacionEstadoCuenta.objects.create(**data)

    def resumen(self, perspective='caja'):
        return resumen_gastos(self.user, self.period, self.period, perspective)

    def test_importe_usd_desconocido_se_conserva_y_no_se_confunde_con_cero(self):
        movement = self.movimiento('10', moneda='USD', monto_clp=Decimal('0'))
        for perspective in ('caja', 'consumo'):
            resumen = self.resumen(perspective)
            line = resumen['categorias'][0]['movimientos'][0]
            self.assertIsNone(line['monto_clp'])
            self.assertIsNone(line['monto_calculo_clp'])
            self.assertEqual(line['metadata']['monto_original'], Decimal('10'))
            self.assertEqual(line['metadata']['moneda_original'], 'USD')
            self.assertFalse(resumen['totales_completos'])
            self.assertEqual(resumen['sin_conversion'], 1)
        self.assertIsNone(movement.monto_abs_clp)
        self.client.force_login(self.user)
        response = self.client.get(reverse('banca_movimientos'), {'direccion': 'EGRESO', 'fecha_desde': '2026-08-01', 'fecha_hasta': '2026-08-31'})
        self.assertContains(response, 'Conversión CLP pendiente')
        self.assertContains(response, 'USD')

    def test_cero_real_y_tasa_bancaria_conocida_si_son_importes_validos(self):
        self.movimiento('0')
        self.movimiento('10', moneda='USD', monto_clp=Decimal('10000'), tipo_cambio_banco=Decimal('1000'))
        resumen = self.resumen()
        self.assertTrue(resumen['totales_completos'])
        self.assertEqual(resumen['totales']['egresos_clp'], Decimal('10000'))

    def test_reimportar_con_tasa_completa_el_importe_y_quita_la_cobertura_pendiente(self):
        movement = self.movimiento('10', moneda='USD', monto_clp=Decimal('0'))
        from bank_sync.models import MovimientoNormalizado
        from banca.services.ingesta import ingestar_movimientos
        result = ingestar_movimientos(self.user, self.bank, [MovimientoNormalizado(
            fecha=movement.fecha, descripcion=movement.descripcion, monto=movement.monto,
            moneda='USD', estado='FACTURADO',
        )], tipo_cambio_usd=Decimal('900'), tipo_producto='CUENTA', nombre_producto=self.account.nombre)
        self.assertEqual(result['nuevos'], 0)
        self.assertEqual(result['ids'], [movement.pk])
        resumen = self.resumen()
        self.assertTrue(resumen['totales_completos'])
        self.assertEqual(resumen['totales']['egresos_clp'], Decimal('9000'))

    def test_devolucion_usd_desconocida_sigue_siendo_egreso_tdc(self):
        self.movimiento('-10', producto=self.card, moneda='USD', monto_clp=Decimal('0'))
        line = self.resumen('consumo')['categorias'][0]['movimientos'][0]
        self.assertEqual(line['direccion'], 'EGRESO')
        self.assertIsNone(line['monto_clp'])
        self.assertIn('pendiente', line['calculo'])

    def test_estado_sin_tasa_y_pago_real_cero_se_distinguen(self):
        estado = self.estado(moneda='USD', monto_facturado_clp_estimado=Decimal('0'))
        self.assertIsNone(monto_facturado_clp(estado)[0])
        resumen = self.resumen()
        self.assertIsNone(resumen['estados_cuenta'][0]['monto_clp'])
        self.assertFalse(resumen['totales_completos'])
        estado.monto_pagado_clp = Decimal('0')
        estado.save()
        self.assertEqual(monto_facturado_clp(estado), (Decimal('0'), False))
        self.assertTrue(self.resumen()['totales_completos'])

    def test_estado_estimado_utiliza_tasa_historica_y_no_una_futura(self):
        estado = self.estado(moneda='USD', monto_facturado=Decimal('10'))
        TipoCambio.objects.create(fecha=date(2026, 7, 1), uf=Decimal('40000'), dolar=Decimal('900'))
        TipoCambio.objects.create(fecha=date(2026, 9, 1), uf=Decimal('40000'), dolar=Decimal('1000'))
        self.assertEqual(monto_facturado_clp(estado), (Decimal('9000'), True))

    def test_api_y_excel_preservan_la_conversion_pendiente_como_ausente(self):
        self.movimiento('10', moneda='USD', monto_clp=Decimal('0'))
        self.client.force_login(self.user)
        query = {'desde': '2026-08', 'hasta': '2026-08'}
        api = self.client.get(reverse('gastos_resumen_api'), query).json()
        self.assertFalse(api['totales_completos'])
        self.assertIsNone(api['categorias'][0]['movimientos'][0]['monto_clp'])
        response = self.client.get(reverse('gastos_exportar_api'), query)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertIsNone(sheet.cell(2, 8).value)
        self.assertIsNone(sheet.cell(2, 9).value)
        self.assertEqual(sheet.cell(2, 12).value, 10)
        self.assertEqual(sheet.cell(2, 13).value, 'USD')
        self.assertEqual(sheet.cell(2, 14).value, 'Pendiente')

    def test_proyeccion_desconocida_no_desaparece_por_tener_clp_cero(self):
        self.movimiento('10', producto=self.card, moneda='USD', monto_clp=Decimal('0'), estado='NO_FACTURADO')
        estado = self.estado(moneda='USD', fecha_vencimiento=date(2026, 9, 10))
        result = proyeccion_no_facturados(self.user, 2026, 8)
        self.assertEqual(result['cantidad'], 1)
        self.assertEqual(result['sin_conversion'], 2)
        self.assertFalse(result['totales_completos'])
        self.assertEqual(result['total_clp'], Decimal('0'))

    def test_producto_excluido_no_suma_movimientos_estados_cobertura_o_presupuesto(self):
        self.account.contabilizar = self.card.contabilizar = False
        self.account.save()
        self.card.save()
        movement = self.movimiento()
        self.estado()
        self.movimiento('500', producto=self.card, estado='NO_FACTURADO')
        self.assertFalse(movement.contabiliza)
        for perspective in ('caja', 'consumo'):
            result = self.resumen(perspective)
            self.assertEqual(result['totales']['egresos_clp'], Decimal('0'))
            self.assertFalse(result['estados_cuenta'])
            self.assertFalse(result['cobertura'])
        self.assertEqual(proyeccion_no_facturados(self.user, 2026, 8)['total_clp'], Decimal('0'))
        self.assertEqual(resumen_presupuesto(self.user, 2026, 8)['ocio_gastado'], Decimal('0'))

    def test_manual_vinculado_a_producto_excluido_se_conserva_como_fuente(self):
        movement = self.movimiento()
        manual = MovimientoManual.objects.create(user=self.user, fecha=movement.fecha, periodo=self.period, direccion='EGRESO', monto=Decimal('100'), monto_clp=Decimal('100'), categoria=self.cat, descripcion='Pago confirmado')
        ConciliacionMovimientoManual.objects.create(movimiento_manual=manual, movimiento_bancario=movement, estado='CONFIRMADA', confirmada_manual=True)
        self.account.contabilizar = False
        self.account.save()
        self.assertEqual(self.resumen()['totales']['egresos_clp'], Decimal('100'))

    def test_promedio_divide_por_meses_vacios_excluye_devoluciones_y_futuros(self):
        self.movimiento('-9000', fecha=date(2026, 8, 10))
        self.movimiento('-9000', producto=self.card, fecha=date(2026, 7, 10))
        self.movimiento('-9000', fecha=date(2026, 9, 10))
        self.assertEqual(promedio_ingresos(self.user, 3, date(2026, 8, 31)), Decimal('3000'))

    def test_presupuesto_incluye_manuales_y_devoluciones_segun_gastos(self):
        MovimientoManual.objects.create(user=self.user, fecha=date(2026, 8, 10), periodo=self.period, direccion='EGRESO', monto=Decimal('1000'), monto_clp=Decimal('1000'), categoria=self.cat, descripcion='Manual')
        self.movimiento('500', producto=self.card)
        self.movimiento('-200', producto=self.card)
        with patch('banca.services.presupuesto.timezone.localdate', return_value=date(2026, 8, 31)):
            presupuesto = resumen_presupuesto(self.user, 2026, 8)
        self.assertEqual(presupuesto['ocio_gastado'], Decimal('1300'))
        self.assertEqual(self.resumen('consumo')['totales']['egresos_clp'], presupuesto['ocio_gastado'])
        filas = detalle_promedios_categoria(self.user, 3, date(2026, 9, 1))
        self.assertEqual(filas[0]['meses'][-1]['monto'], Decimal('1300'))

    def test_reportes_no_usan_movimientos_ni_conversiones_de_otro_usuario(self):
        foreign = Producto.objects.create(user=self.other, banco=self.bank, nombre='Ajena', tipo='CUENTA')
        self.movimiento('-9999', user=self.other, producto=foreign)
        self.movimiento('10', user=self.other, producto=foreign, moneda='USD', monto_clp=Decimal('0'))
        self.assertTrue(self.resumen()['totales_completos'])
        self.assertEqual(self.resumen()['totales']['ingresos_clp'], Decimal('0'))
        self.assertEqual(promedio_ingresos(self.user, 3, date(2026, 8, 31)), Decimal('0'))

    def test_dashboard_clasifica_cada_linea_una_vez_aunque_compartan_categoria(self):
        subs = CategoriaTematica.objects.create(user=self.user, nombre='Suscripciones', aplica_a='EGRESO', esencialidad_defecto='FIJO')
        self.movimiento('200', categoria=subs, esencialidad='FIJO')
        self.movimiento('100', esencialidad='FIJO')
        self.movimiento('300', esencialidad='OCIO')
        self.client.force_login(self.user)
        response = self.client.get(reverse('dashboard'), {'year': '2026', 'month': '8'})
        self.assertEqual(response.context['gastos_fijos'], Decimal('100'))
        self.assertEqual(response.context['suscripciones'], Decimal('200'))
        self.assertEqual(response.context['otros'], Decimal('300'))
        self.assertEqual(response.context['gastos_mes'], Decimal('600'))

    def test_saldos_excluidos_no_suman_cupos_ni_deudas(self):
        SaldoSnapshot.objects.create(user=self.user, producto=self.card, fecha=date(2026, 8, 1), deuda_facturada_clp=Decimal('500'), cupo_total_clp=Decimal('1000'))
        self.card.contabilizar = False
        self.card.save()
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse('banca_deudas')).context['total_deuda'], Decimal('0'))
        from banca.services.analisis import cupos_tdc, monto_a_pagar_tdc
        self.assertEqual(cupos_tdc(self.user)['total_clp'], Decimal('0'))
        self.assertFalse(monto_a_pagar_tdc(self.user))
