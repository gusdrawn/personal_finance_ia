"""Importación, ingesta y conciliación de Houm/Neat sin duplicar la banca."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone as datetime_timezone
from decimal import Decimal, InvalidOperation
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from django.db import transaction
from django.db.models import Q
from django.utils import timezone

from bank_sync.exceptions import AuthError
from bank_sync.models import (
    DetalleLiquidacionNormalizado,
    LiquidacionArriendoNormalizada,
    PagoExternoNormalizado,
)
from banca.models import (
    CategoriaTematica,
    ConciliacionEventoExterno,
    CredencialBanco,
    DetalleLiquidacionHoum,
    EventoFinancieroExterno,
    ImportacionFuenteExterna,
    LiquidacionHoum,
    MovimientoBancario,
    PagoNeat,
    SyncLog,
    VinculoPropiedadExterna,
)


MAX_FILE_SIZE = 15 * 1024 * 1024


def _key(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch)).lower()
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def _decimal(value) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace("$", "").replace("CLP", "").replace(" ", "")
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        tail = text.rsplit(",", 1)[-1]
        text = text.replace(".", "")
        text = text.replace(",", "." if len(tail) <= 2 else "")
    elif text.count(".") > 1 or ("." in text and len(text.rsplit(".", 1)[-1]) == 3):
        text = text.replace(".", "")
    try:
        amount = Decimal(text)
    except InvalidOperation:
        return Decimal("0")
    return -amount if negative else amount


def _date(value, *, period=False):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    else:
        text = str(value).strip()
        parsed = None
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%m/%Y", "%Y-%m"):
            try:
                parsed = datetime.strptime(text[:10], fmt).date()
                break
            except ValueError:
                continue
        if parsed is None:
            return None
    return parsed.replace(day=1) if period else parsed


def _first(row: dict, *names, default=""):
    for name in names:
        value = row.get(name)
        if value not in (None, ""):
            return value
    return default


def _serializable_row(row: dict) -> dict:
    result = {}
    for key, value in row.items():
        if isinstance(value, (date, datetime)):
            value = value.isoformat()
        elif isinstance(value, Decimal):
            value = str(value)
        result[str(key)] = value
    return result


def _proof_row(text: str) -> dict:
    """Extrae campos frecuentes de un comprobante/correo Neat sin depender del HTML."""
    plain = re.sub(r"<[^>]+>", " ", text or "")
    plain = re.sub(r"\s+", " ", plain).strip()
    patterns = {
        "id_pago": r"(?:id|identificador)(?:\s+del)?\s+pago\s*[:#-]?\s*([A-Z0-9_-]+)",
        "estado": r"estado\s*[:#-]?\s*([A-Za-zÁÉÍÓÚáéíóú_ ]+?)(?=\s+(?:monto|fecha|total|comisi[oó]n|destinatario)\b|$)",
        "monto": r"(?:monto(?:\s+pagado)?|importe)\s*[:#-]?\s*\$?\s*([0-9.,]+)",
        "comision": r"comisi[oó]n\s*[:#-]?\s*\$?\s*([0-9.,]+)",
        "total_cargado": r"total(?:\s+cargado)?\s*[:#-]?\s*\$?\s*([0-9.,]+)",
        "fecha_cargo": r"fecha(?:\s+de)?\s+cargo\s*[:#-]?\s*([0-9/-]{8,10})",
        "fecha": r"fecha(?:\s+de)?\s+pago\s*[:#-]?\s*([0-9/-]{8,10})",
        "destinatario": r"(?:destinatario|beneficiario)\s*[:#-]?\s*(.+?)(?=\s+(?:monto|fecha|total|comisi[oó]n|estado|tarjeta)\b|$)",
        "tarjeta_ultimos4": r"(?:tarjeta|terminada en|[uú]ltimos 4)\D{0,20}([0-9]{4})\b",
    }
    row = {}
    for key, pattern in patterns.items():
        match = re.search(pattern, plain, flags=re.IGNORECASE)
        if match:
            row[key] = match.group(1).strip()
    if not row.get("monto") and not row.get("total_cargado"):
        raise ValueError("El comprobante no contiene un monto reconocible.")
    row["texto_origen"] = plain[:5000]
    return row


def _tabular_rows(content: bytes, name: str) -> list[dict]:
    lower = name.lower()
    if lower.endswith((".txt", ".eml", ".html", ".htm")):
        return [_proof_row(content.decode("utf-8-sig", errors="replace"))]
    if lower.endswith(".pdf"):
        import pdfplumber

        with pdfplumber.open(io.BytesIO(content)) as document:
            text = "\n".join(page.extract_text() or "" for page in document.pages)
        return [_proof_row(text)]
    if lower.endswith(".json"):
        payload = json.loads(content.decode("utf-8-sig"))
        if isinstance(payload, list):
            rows = payload
        else:
            rows = payload.get("data") or payload.get("payments") or payload.get("settlements") or []
        return [{_key(k): v for k, v in row.items()} for row in rows if isinstance(row, dict)]
    if lower.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [
            {_key(k): v for k, v in row.items()}
            for row in csv.DictReader(io.StringIO(text), dialect=dialect)
        ]
    if lower.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        matrix = list(sheet.iter_rows(values_only=True))
        header_index = None
        headers = []
        expected = {"monto", "renta", "arriendo", "deposito", "propiedad", "pago", "id"}
        for index, values in enumerate(matrix[:25]):
            candidate = [_key(v) for v in values]
            if sum(any(word in cell for word in expected) for cell in candidate) >= 2:
                header_index, headers = index, candidate
                break
        if header_index is None:
            raise ValueError("No se encontró una fila de encabezados reconocible.")
        rows = []
        for values in matrix[header_index + 1 :]:
            row = {headers[i]: value for i, value in enumerate(values) if i < len(headers) and headers[i]}
            if any(value not in (None, "") for value in row.values()):
                rows.append(row)
        return rows
    raise ValueError("Formato no soportado; usa XLSX, CSV o JSON.")


def parsear_neat(content: bytes, name: str) -> list[PagoExternoNormalizado]:
    rows = _tabular_rows(content, name)
    pagos = []
    for index, row in enumerate(rows, start=1):
        amount = _decimal(_first(row, "monto", "importe", "amount", "monto_pagado", "valor"))
        fee = _decimal(_first(row, "comision", "fee", "commission"))
        total = _decimal(_first(row, "total_cargado", "total", "cargo_total"))
        if not amount and total:
            amount = total - fee
        external_id = str(_first(row, "id_pago", "identificador_de_pago", "payment_id", "id", default="")).strip()
        event_date = _date(_first(row, "fecha", "fecha_pago", "dia_de_pago", "created_at"))
        payment_date = _date(_first(row, "fecha_cargo", "charge_date", "dia_de_pago")) or event_date
        deposit_date = _date(_first(row, "fecha_transferencia", "transfer_date", "dia_de_deposito"))
        if not external_id:
            seed = f"{event_date}|{amount}|{_first(row, 'empresa', 'destinatario', 'cuenta')}|{index}"
            external_id = "archivo-" + hashlib.sha256(seed.encode()).hexdigest()[:24]
        status = _key(_first(row, "estado", "status", default="en_proceso"))
        status_map = {
            "pagado": "CONFIRMADO", "exitoso": "CONFIRMADO", "realizado": "CONFIRMADO",
            "en_proceso": "EN_PROCESO", "pendiente": "EN_PROCESO",
            "rechazado": "FALLIDO", "fallido": "FALLIDO", "anulado": "REVERSADO",
        }
        pagos.append(
            PagoExternoNormalizado(
                external_id=external_id,
                estado=status_map.get(status, "EN_PROCESO"),
                monto_pagado=amount,
                comision=fee,
                descuentos=_decimal(_first(row, "descuento", "descuentos")),
                cuenta=str(_first(row, "cuenta", "empresa", "account")),
                tipo_cuenta=str(_first(row, "tipo_cuenta", "tipo", "account_type")),
                destinatario=str(_first(row, "destinatario", "nombre_destinatario", "beneficiario", "recipient")),
                tarjeta_ultimos4=str(_first(row, "tarjeta_ultimos4", "ultimos_4", "card_last4"))[-4:],
                fecha_creacion=event_date,
                fecha_cargo=payment_date,
                fecha_transferencia_esperada=_date(_first(row, "fecha_esperada", "expected_transfer_date")),
                fecha_transferencia=deposit_date,
                comprobante_referencia=str(_first(row, "comprobante", "receipt", "receipt_id")),
                metadata=_serializable_row(row),
            )
        )
    return pagos


def parsear_houm(content: bytes, name: str) -> list[LiquidacionArriendoNormalizada]:
    rows = _tabular_rows(content, name)
    liquidaciones = []
    for index, row in enumerate(rows, start=1):
        property_id = str(_first(row, "id_propiedad", "property_id", "propiedad_id", "codigo", default="")).strip()
        property_name = str(_first(row, "propiedad", "property_name", "direccion", "inmueble", default=property_id)).strip()
        period = _date(_first(row, "periodo", "mes", "fecha_periodo"), period=True)
        deposit_date = _date(_first(
            row, "fecha_deposito", "fecha_pago", "fecha_de_pago", "deposit_date", "fecha",
        ))
        if period is None and deposit_date:
            period = deposit_date.replace(day=1)
        if not property_id:
            property_id = _key(property_name)
        if not property_id or not period:
            continue
        gross = _decimal(_first(row, "renta_bruta", "arriendo", "renta", "gross_rent", "ingreso_bruto"))
        net = _decimal(_first(
            row, "deposito_neto", "deposito", "deposito_a_propietario",
            "arriendo_depositado", "neto", "net_deposit", "ingreso_neto",
        ))
        commission = abs(_decimal(_first(
            row, "plan_houm", "comision_houm", "comision", "commission",
            "administracion", "comision_administracion", "costo_administracion",
        )))
        vat = abs(_decimal(_first(row, "iva", "vat")))
        brokerage = abs(_decimal(_first(row, "corretaje", "brokerage")))
        promotions = abs(_decimal(_first(row, "promociones", "descuentos", "promotions")))
        repairs = abs(_decimal(_first(row, "reparaciones", "repairs")))
        services = abs(_decimal(_first(row, "servicios", "servicios_basicos", "services")))
        adjustments = _decimal(_first(row, "ajustes", "otros", "adjustments"))
        if not net:
            net = gross - commission - vat - brokerage - promotions - repairs - services + adjustments
        external_id = str(_first(row, "id_liquidacion", "settlement_id", "id_pago", "id", default="")).strip()
        if not external_id:
            # El reporte mensual oficial no trae ID de liquidación. La propiedad y el
            # período forman una identidad estable para que un archivo corregido actualice
            # la liquidación en vez de crear una segunda versión del mismo mes.
            seed = f"houm|{property_id}|{period.isoformat()}"
            external_id = "archivo-" + hashlib.sha256(seed.encode()).hexdigest()[:24]
        details = []
        for kind, label, amount in (
            ("COMISION", "Comisión Houm", commission), ("IVA", "IVA", vat),
            ("CORRETAJE", "Corretaje", brokerage), ("PROMOCION", "Promociones y descuentos", promotions),
            ("REPARACION", "Reparaciones", repairs), ("SERVICIO", "Servicios básicos", services),
            ("AJUSTE", "Otros ajustes", adjustments),
        ):
            if amount:
                details.append(DetalleLiquidacionNormalizado(label, amount, kind))
        liquidaciones.append(
            LiquidacionArriendoNormalizada(
                external_id=external_id,
                propiedad_id=property_id,
                propiedad_nombre=property_name or property_id,
                periodo=period,
                renta_bruta=gross,
                deposito_neto=net,
                estado="CONFIRMADO" if deposit_date else "EN_PROCESO",
                fecha_deposito=deposit_date,
                comision=commission,
                iva=vat,
                corretaje=brokerage,
                promociones=promotions,
                reparaciones=repairs,
                servicios=services,
                ajustes=adjustments,
                detalles=details,
                metadata=_serializable_row(row),
            )
        )
    return liquidaciones


def _credential(user, source):
    return CredencialBanco.objects.filter(user=user, conector=source).first()


NEAT_CATEGORIAS_CONOCIDAS = (
    ("Servicios básicos", ("nuevo sur", "aguas andinas", "metrogas", "luz linares", "enel")),
    ("Telefonía / Internet", ("entel", "movistar", "vtr", "wom", "claro")),
    ("Financiamiento / Hipoteca", ("hipotecario",)),
    ("Arriendo / Hipoteca", ("arriendo",)),
    ("Financiamiento / Cuotas", ("credito", "crédito")),
)


def _categoria_conocida_neat(user, glosa):
    """Resuelve servicios reconocibles sin depender de disponibilidad de la IA."""
    for nombre, patrones in NEAT_CATEGORIAS_CONOCIDAS:
        if any(patron in glosa for patron in patrones):
            categoria, _ = CategoriaTematica.objects.get_or_create(
                user=user,
                nombre=nombre,
                defaults={"aplica_a": "EGRESO", "origen": "SISTEMA", "esencialidad_defecto": "FIJO"},
            )
            return categoria
    return None


def _clasificar_evento_neat(user, event, payment, *, usar_ia=True):
    """Aplica primero reglas aprendidas y luego la misma IA usada por la banca."""
    if event.categoria_es_manual:
        return
    from banca.services.clasificacion_ia import (
        _categoria,
        _clasificador_usuario,
        guardar_mapeo_comercio,
        normalizar_glosa,
    )
    from banca.models import MapeoComercio

    alias = str((event.metadata or {}).get("alias") or "").strip()
    texto = " ".join(
        parte for parte in (payment.destinatario, payment.cuenta, alias) if parte
    ) or "NeatPagos"
    glosa = normalizar_glosa(texto)
    mapeo = next(
        (item for item in MapeoComercio.objects.filter(user=user).select_related("categoria") if item.patron in glosa),
        None,
    )
    if mapeo:
        event.categoria = mapeo.categoria
        event.fuente_clasificacion = "REGLA"
        event.confianza_ia = None
        event.save(update_fields=["categoria", "fuente_clasificacion", "confianza_ia", "updated_at"])
        return
    categoria_conocida = _categoria_conocida_neat(user, glosa)
    if categoria_conocida:
        event.categoria = categoria_conocida
        event.fuente_clasificacion = "REGLA"
        event.confianza_ia = None
        event.save(update_fields=["categoria", "fuente_clasificacion", "confianza_ia", "updated_at"])
        return
    if not usar_ia:
        return
    categorias = list(CategoriaTematica.objects.filter(user=user, activa=True).values_list("nombre", flat=True))
    resultado = _clasificador_usuario(user).clasificar([glosa], categorias)
    if not resultado:
        return
    nombre = str(resultado[0].get("categoria", "")).strip()
    if not nombre:
        return
    categoria = _categoria(user, nombre)
    if categoria.origen == "SISTEMA" and nombre not in categorias:
        categoria.origen = "IA"
        categoria.save(update_fields=["origen"])
    event.categoria = categoria
    event.fuente_clasificacion = "IA"
    event.confianza_ia = resultado[0].get("confianza")
    event.save(update_fields=["categoria", "fuente_clasificacion", "confianza_ia", "updated_at"])
    guardar_mapeo_comercio(user, glosa, categoria, origen="IA")


def importar_archivo_externo(user, source: str, uploaded_file):
    if source not in {"neat", "houm"}:
        raise ValueError("Fuente externa no soportada.")
    if uploaded_file.size > MAX_FILE_SIZE:
        raise ValueError("El archivo supera el límite de 15 MB.")
    content = uploaded_file.read()
    digest = hashlib.sha256(content).hexdigest()
    existing = ImportacionFuenteExterna.objects.filter(
        user=user, fuente=source, hash_archivo=digest
    ).first()
    if existing:
        return existing, True
    record = ImportacionFuenteExterna.objects.create(
        user=user,
        credencial=_credential(user, source),
        fuente=source,
        nombre_archivo=uploaded_file.name or "archivo",
        mime_type=getattr(uploaded_file, "content_type", "") or "application/octet-stream",
        contenido_privado=content,
        hash_archivo=digest,
    )
    try:
        parsed = parsear_neat(content, record.nombre_archivo) if source == "neat" else parsear_houm(content, record.nombre_archivo)
        if not parsed:
            raise ValueError("El archivo no contiene registros reconocibles para esta fuente.")
        result = ingestar_pagos_neat(user, parsed, record) if source == "neat" else ingestar_liquidaciones_houm(user, parsed, record)
        record.elementos_detectados = len(parsed)
        record.nuevos = result["nuevos"]
        record.actualizados = result["actualizados"]
        periods = [getattr(item, "periodo", None) for item in parsed if getattr(item, "periodo", None)]
        record.periodo = min(periods) if periods else None
        record.mensaje = "OK"
    except Exception as exc:
        record.estado = "ERROR"
        record.mensaje = f"{type(exc).__name__}: {exc}"
    record.save()
    if record.estado == "PROCESADO":
        conciliar_eventos(user)
    return record, False


def _find_property_link(user, property_id, property_name):
    link = VinculoPropiedadExterna.objects.filter(
        user=user, fuente="houm", external_property_id=property_id
    ).select_related("departamento").first()
    if link:
        return link
    from departamentos.models import Departamento

    property_key = _key(property_name)
    property_words = property_key.replace("_", " ")
    property_codes = {property_key.replace("_", "")}
    # Houm escribe la unidad como 107F o "107 E", mientras el proyecto usa F107/E107.
    for number, tower in re.findall(r"\b(\d{1,4})\s*([a-z])\b", property_words):
        property_codes.add(f"{tower}{number}")
    for tower, number in re.findall(r"\b([a-z])\s*(\d{1,4})\b", property_words):
        property_codes.add(f"{tower}{number}")
    candidates = [
        department
        for department in Departamento.objects.filter(user=user, activo=True)
        if _key(department.codigo).replace("_", "") in property_codes
        or _key(department.codigo) in property_key
    ]
    if len(candidates) == 1:
        return VinculoPropiedadExterna.objects.create(
            user=user, fuente="houm", external_property_id=property_id,
            nombre_externo=property_name, departamento=candidates[0]
        )
    return None


@transaction.atomic
def vincular_propiedad(user, property_id, departamento):
    """Confirma un mapeo Houm y reprocesa sus liquidaciones pendientes."""
    if departamento.user_id != user.id:
        raise ValueError("El departamento no pertenece al usuario.")
    pending = LiquidacionHoum.objects.filter(
        evento__user=user, external_property_id=property_id
    ).select_related("evento")
    name = pending.values_list("propiedad_nombre", flat=True).first() or property_id
    link, _ = VinculoPropiedadExterna.objects.update_or_create(
        user=user, fuente="houm", external_property_id=property_id,
        defaults={"nombre_externo": name, "departamento": departamento, "activo": True},
    )
    for liquidation in pending:
        liquidation.vinculo_propiedad = link
        liquidation.save(update_fields=["vinculo_propiedad"])
        event = liquidation.evento
        event.requiere_revision = False
        if event.estado == "REQUIERE_REVISION":
            event.estado = "CONFIRMADO" if event.fecha_confirmacion else "EN_PROCESO"
        event.save(update_fields=["requiere_revision", "estado", "updated_at"])
    return link


@transaction.atomic
def ingestar_pagos_neat(user, payments, import_record=None, credential=None):
    created_count = updated_count = 0
    for payment in payments:
        defaults = {
            "credencial": credential or _credential(user, "neat"),
            "importacion": import_record,
            "tipo": "PAGO", "estado": payment.estado,
            "periodo": (payment.fecha_cargo or payment.fecha_creacion).replace(day=1) if (payment.fecha_cargo or payment.fecha_creacion) else None,
            "fecha_evento": payment.fecha_creacion, "fecha_cargo": payment.fecha_cargo,
            "fecha_esperada": payment.fecha_transferencia_esperada,
            "fecha_confirmacion": payment.fecha_transferencia,
            "moneda": payment.moneda, "monto_bruto": payment.monto_pagado,
            "comision": payment.comision, "descuentos": payment.descuentos,
            "monto_neto": payment.total_cargado, "metadata": payment.metadata,
        }
        event, created = EventoFinancieroExterno.objects.update_or_create(
            user=user, fuente="neat", external_id=payment.external_id, defaults=defaults
        )
        PagoNeat.objects.update_or_create(
            evento=event,
            defaults={"cuenta": payment.cuenta, "tipo_cuenta": payment.tipo_cuenta,
                      "destinatario": payment.destinatario, "tarjeta_ultimos4": payment.tarjeta_ultimos4,
                      "monto_pagado": payment.monto_pagado, "total_cargado": payment.total_cargado,
                      "comprobante_referencia": payment.comprobante_referencia},
        )
        _clasificar_evento_neat(user, event, payment)
        created_count += int(created)
        updated_count += int(not created)
    return {"nuevos": created_count, "actualizados": updated_count}


@transaction.atomic
def ingestar_liquidaciones_houm(user, settlements, import_record=None, credential=None):
    created_count = updated_count = 0
    categoria_houm, _ = CategoriaTematica.objects.get_or_create(
        user=user,
        nombre="Arriendos Houm",
        defaults={"aplica_a": "INGRESO", "origen": "SISTEMA", "esencialidad_defecto": "FIJO"},
    )
    for settlement in settlements:
        link = _find_property_link(user, settlement.propiedad_id, settlement.propiedad_nombre)
        event, created = EventoFinancieroExterno.objects.update_or_create(
            user=user, fuente="houm", external_id=settlement.external_id,
            defaults={"credencial": credential or _credential(user, "houm"), "importacion": import_record,
                      "tipo": "LIQUIDACION", "estado": settlement.estado if link else "REQUIERE_REVISION", "periodo": settlement.periodo,
                      "fecha_evento": settlement.fecha_deposito or settlement.periodo,
                      "fecha_esperada": settlement.periodo.replace(day=6),
                      "fecha_confirmacion": settlement.fecha_deposito, "moneda": settlement.moneda,
                      "monto_bruto": settlement.renta_bruta, "comision": settlement.comision + settlement.iva,
                      "descuentos": settlement.corretaje + settlement.promociones + settlement.reparaciones + settlement.servicios,
                      "monto_neto": settlement.deposito_neto, "metadata": settlement.metadata,
                      "categoria": categoria_houm, "fuente_clasificacion": "REGLA",
                      "requiere_revision": link is None},
        )
        liquidation, _ = LiquidacionHoum.objects.update_or_create(
            evento=event,
            defaults={"vinculo_propiedad": link, "external_property_id": settlement.propiedad_id,
                      "propiedad_nombre": settlement.propiedad_nombre, "renta_bruta": settlement.renta_bruta,
                      "iva": settlement.iva, "corretaje": settlement.corretaje,
                      "promociones": settlement.promociones, "reparaciones": settlement.reparaciones,
                      "servicios": settlement.servicios, "ajustes": settlement.ajustes,
                      "deposito_neto": settlement.deposito_neto},
        )
        liquidation.detalles.all().delete()
        DetalleLiquidacionHoum.objects.bulk_create([
            DetalleLiquidacionHoum(liquidacion=liquidation, tipo=item.tipo, concepto=item.concepto,
                                   monto=item.monto, metadata=item.metadata)
            for item in settlement.detalles
        ])
        created_count += int(created)
        updated_count += int(not created)
    return {"nuevos": created_count, "actualizados": updated_count}




def _candidate_movements(event, amount, start, end, *, charge):
    qs = MovimientoBancario.objects.filter(
        user=event.user, fecha__range=(start, end)
    ).exclude(estado="DESCARTADO_CIERRE")
    qs = qs.filter(monto_clp__gte=0) if charge else qs.filter(monto_clp__lt=0)
    return [m for m in qs if abs(abs(m.monto_clp) - abs(amount)) <= Decimal("1")]


def _banking_window(reference, days):
    start = end = reference
    remaining = days
    while remaining:
        start -= timedelta(days=1)
        if start.weekday() < 5:
            remaining -= 1
    remaining = days
    while remaining:
        end += timedelta(days=1)
        if end.weekday() < 5:
            remaining -= 1
    return start, end


def _sincronizar_categoria_movimiento_neat(movimiento):
    """Traslada una categoría Neat inequívoca sin pisar decisiones manuales."""
    if movimiento.categoria_es_manual:
        return False
    categorias = {
        conciliacion.evento.categoria
        for conciliacion in movimiento.conciliaciones_externas.all()
        if conciliacion.evento.fuente == "neat" and conciliacion.evento.categoria_id
    }
    if len(categorias) != 1:
        return False
    categoria = categorias.pop()
    if movimiento.categoria_id == categoria.pk and movimiento.fuente_clasificacion == "REGLA":
        return False
    movimiento.categoria = categoria
    movimiento.esencialidad = categoria.esencialidad_defecto
    movimiento.fuente_clasificacion = "REGLA"
    movimiento.confianza_ia = None
    movimiento.clasificado_at = timezone.now()
    movimiento.save(update_fields=[
        "categoria", "esencialidad", "fuente_clasificacion", "confianza_ia",
        "clasificado_at", "updated_at",
    ])
    return True


@transaction.atomic
def conciliar_eventos(user):
    matched = review = classified = 0
    events = EventoFinancieroExterno.objects.filter(user=user).exclude(estado="FALLIDO")
    for event in events.select_related("pago_neat", "liquidacion_houm"):
        if event.fuente == "neat" and not event.categoria_es_manual:
            categoria_anterior = event.categoria_id
            _clasificar_evento_neat(user, event, event.pago_neat, usar_ia=False)
            classified += int(event.categoria_id != categoria_anterior)
        if event.conciliaciones.exists():
            continue
        reference = event.fecha_cargo or event.fecha_confirmacion or event.fecha_evento
        if not reference:
            continue
        charge = event.fuente == "neat"
        window = 3 if charge else 5
        start, end = _banking_window(reference, window)
        candidates = _candidate_movements(event, event.monto_neto, start, end, charge=charge)
        method = "MONTO_FECHA"
        identified = [m for m in candidates if event.external_id.lower() in (m.descripcion or "").lower()]
        if identified:
            candidates, method = identified, "ID"
        if charge:
            last4 = getattr(event.pago_neat, "tarjeta_ultimos4", "")
            neat_candidates = [m for m in candidates if "NEAT" in (m.descripcion or "").upper()]
            card_candidates = [m for m in candidates if last4 and last4 in (m.descripcion or "")]
            candidates = identified or card_candidates or neat_candidates
        if len(candidates) == 1:
            ConciliacionEventoExterno.objects.create(
                evento=event, movimiento=candidates[0], rol="CARGO" if charge else "DEPOSITO",
                monto_asignado=event.monto_neto, confianza=Decimal("0.980" if method == "ID" else "0.950"), metodo=method
            )
            matched += 1
        elif len(candidates) > 1:
            event.requiere_revision = True
            event.estado = "REQUIERE_REVISION"
            event.save(update_fields=["requiere_revision", "estado", "updated_at"])
            review += 1

    # Un cargo Neat puede agrupar pagos divididos entre varios destinatarios.
    pending_neat = list(
        events.filter(fuente="neat", conciliaciones__isnull=True)
        .select_related("pago_neat")
        .distinct()
    )
    by_charge = defaultdict(list)
    for event in pending_neat:
        if event.fecha_cargo:
            by_charge[(event.fecha_cargo, event.pago_neat.tarjeta_ultimos4)].append(event)
    for (charge_date, last4), group in by_charge.items():
        if len(group) < 2:
            continue
        total = sum((event.monto_neto for event in group), Decimal("0"))
        start, end = _banking_window(charge_date, 3)
        candidates = _candidate_movements(group[0], total, start, end, charge=True)
        candidates = [
            movement for movement in candidates
            if "NEAT" in (movement.descripcion or "").upper()
            or (last4 and last4 in (movement.descripcion or ""))
        ]
        if len(candidates) == 1:
            for event in group:
                ConciliacionEventoExterno.objects.create(
                    evento=event, movimiento=candidates[0], rol="CARGO",
                    monto_asignado=event.monto_neto, confianza=Decimal("0.900"),
                    metodo="MONTO_FECHA",
                )
                matched += 1

    # Un depósito bancario puede agrupar varias propiedades del mismo período.
    pending_houm = list(events.filter(fuente="houm", conciliaciones__isnull=True).distinct())
    by_period = defaultdict(list)
    for event in pending_houm:
        if event.periodo:
            by_period[event.periodo].append(event)
    for period, group in by_period.items():
        if len(group) < 2:
            continue
        total = sum((event.monto_neto for event in group), Decimal("0"))
        start, end = period.replace(day=1), period.replace(day=10)
        candidates = _candidate_movements(group[0], total, start, end, charge=False)
        if len(candidates) == 1:
            for event in group:
                ConciliacionEventoExterno.objects.create(
                    evento=event, movimiento=candidates[0], rol="DEPOSITO",
                    monto_asignado=event.monto_neto, confianza=Decimal("0.900"), metodo="MONTO_FECHA"
                )
                matched += 1

    movimientos_neat = (
        MovimientoBancario.objects.filter(
            user=user,
            conciliaciones_externas__evento__fuente="neat",
        )
        .prefetch_related(
            "conciliaciones_externas__evento__categoria",
        )
        .distinct()
    )
    for movimiento in movimientos_neat:
        classified += int(_sincronizar_categoria_movimiento_neat(movimiento))
    return {"conciliados": matched, "revision": review, "clasificados": classified}


def periodo_actual(fecha=None):
    fecha = fecha or timezone.localdate()
    return fecha.replace(day=1)


def _local_now(credential, now=None):
    if now is not None:
        return now
    try:
        zone = ZoneInfo(credential.zona_horaria)
    except (ZoneInfoNotFoundError, ValueError):
        zone = ZoneInfo("America/Santiago")
    return timezone.now().astimezone(zone)


def calcular_proxima_sincronizacion(credential, now=None):
    """Primera hora válida futura de la ventana mensual configurada."""
    if not credential.activo or not credential.sincronizacion_automatica:
        return None
    now = _local_now(credential, now)
    for offset in range(370):
        candidate_day = now.date() + timedelta(days=offset)
        if not (credential.dia_inicio <= candidate_day.day <= credential.dia_fin):
            continue
        candidate = datetime.combine(candidate_day, credential.hora_sincronizacion, now.tzinfo)
        if candidate <= now:
            continue
        if credential.conector == "houm" and credential.detener_al_completar:
            if houm_completo(credential, candidate_day.replace(day=1)):
                continue
        return candidate.astimezone(datetime_timezone.utc)
    return None


def houm_completo(credential, period=None):
    period = period or periodo_actual()
    links = VinculoPropiedadExterna.objects.filter(user=credential.user, fuente="houm", activo=True)
    if not links.exists():
        return False
    captured = LiquidacionHoum.objects.filter(
        vinculo_propiedad__in=links, evento__periodo=period
    ).values_list("vinculo_propiedad_id", flat=True).distinct()
    return set(captured) == set(links.values_list("id", flat=True))


def debe_sincronizar(credential, now=None):
    now = _local_now(credential, now)
    today = now.date()
    if not credential.aplica_en_fecha(today):
        return False
    if now.time().replace(tzinfo=None) < credential.hora_sincronizacion:
        return False
    period = periodo_actual(today)
    if credential.detener_al_completar and credential.conector == "houm" and houm_completo(credential, period):
        return False
    last_sync = None
    if credential.ultima_sincronizacion:
        last_sync = credential.ultima_sincronizacion.astimezone(now.tzinfo).date()
    return last_sync != today


def run_external_sync(credential, *, manual=False, period=None):
    """Sincronización read-only con lock persistente; importación queda como fallback."""
    from bank_sync.registry import get_connector

    period = period or periodo_actual()
    stale = timezone.now() - timedelta(hours=1)
    CredencialBanco.objects.filter(pk=credential.pk, sincronizacion_en_curso=True, sincronizacion_iniciada__lt=stale).update(
        sincronizacion_en_curso=False
    )
    acquired = CredencialBanco.objects.filter(pk=credential.pk, sincronizacion_en_curso=False).update(
        sincronizacion_en_curso=True, sincronizacion_iniciada=timezone.now()
    )
    if not acquired:
        return None
    credential.refresh_from_db()
    log = SyncLog.objects.create(user=credential.user, banco=credential.banco, modo="externo")
    connector = None
    try:
        if not manual and not debe_sincronizar(credential):
            log.exito, log.mensaje = True, "Fuera de ventana o período ya completo"
            return log
        if not credential.conector_experimental or not credential.base_url:
            raise AuthError("Conector automático no habilitado; usa la importación de archivos.")
        connector = get_connector(credential.conector, base_url=credential.base_url)
        connector.authenticate(api_key=credential.get_api_key(), access_token=credential.get_sesion())
        if credential.conector == "neat":
            result = connector.get_payments(desde=period, cursor=credential.ultimo_cursor)
            counts = ingestar_pagos_neat(credential.user, result.elementos, credential=credential)
        else:
            result = connector.get_settlements(periodo=period, cursor=credential.ultimo_cursor)
            counts = ingestar_liquidaciones_houm(credential.user, result.elementos, credential=credential)
        conciliar_eventos(credential.user)
        credential.ultimo_cursor = result.cursor
        credential.ultimo_periodo = period
        credential.ultima_sincronizacion = timezone.now()
        credential.errores_consecutivos = 0
        credential.requiere_reautenticacion = False
        credential.ultimo_mensaje = f"OK: {len(result.elementos)} elemento(s)"
        local_today = _local_now(credential).date()
        if credential.conector == "houm" and local_today.day >= credential.dia_fin and not houm_completo(credential, period):
            pending = VinculoPropiedadExterna.objects.filter(
                user=credential.user, fuente="houm", activo=True
            ).exclude(liquidaciones__evento__periodo=period).count()
            credential.ultimo_mensaje = f"Atención: faltan {pending} propiedad(es) al cerrar la ventana."
        log.exito, log.nuevos, log.actualizados, log.mensaje = True, counts["nuevos"], counts["actualizados"], "OK"
    except Exception as exc:
        credential.errores_consecutivos += 1
        credential.requiere_reautenticacion = isinstance(exc, AuthError)
        credential.ultimo_mensaje = f"{type(exc).__name__}: {exc}"
        credential.ultima_sincronizacion = timezone.now()
        log.exito, log.mensaje = False, credential.ultimo_mensaje
    finally:
        if connector is not None:
            try:
                connector.close()
            except Exception:
                pass
        credential.proxima_sincronizacion = calcular_proxima_sincronizacion(credential)
        credential.sincronizacion_en_curso = False
        credential.sincronizacion_iniciada = None
        credential.save()
        log.finalizado = timezone.now()
        log.save()
    return log


def probar_conexion(credential):
    """Comprueba autenticación y lectura sin persistir información financiera."""
    from bank_sync.registry import get_connector

    if not credential.conector_experimental or not credential.base_url:
        raise AuthError("Activa el endpoint experimental y configura una URL autorizada.")
    connector = get_connector(credential.conector, base_url=credential.base_url)
    try:
        connector.authenticate(
            api_key=credential.get_api_key(), access_token=credential.get_sesion()
        )
        if credential.conector == "houm":
            result = connector.get_properties()
            return len(result)
        else:
            result = connector.get_payments(desde=timezone.localdate(), cursor="")
            return len(result.elementos)
    finally:
        connector.close()
